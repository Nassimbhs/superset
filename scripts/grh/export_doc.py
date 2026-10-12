"""Generate the client Word document describing every GRH dashboard from its spec:
dashboards summary with public URLs, calculation rules in plain language, one chart
table and one filter table per dashboard. No database detail (tables, columns, SQL).

Usage: python export_doc.py [--out exports/Documentation_Dashboards_GRH.docx]
                            [--base-url https://superset.client] [--cod-soc 01]
Requires: pip install python-docx
"""
from __future__ import annotations

import argparse
import importlib
import json
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from grh_builder import HERE, SUITE, TimeFilter

DEFAULT_OUT = HERE / "exports" / "Documentation_Dashboards_GRH.docx"

DESCRIPTIONS = {
    "synthese": "Vue mensuelle de synthèse : effectif payé, entrées/sorties, turnover, rétention, encadrement, "
                "absentéisme toutes causes et masse salariale rapportée à l'effectif.",
    "effectif": "Photographie de l'effectif actif : répartition par direction, sexe, âge, ancienneté, catégorie, "
                "grade, région, fonction et mouvements annuels.",
    "conges": "Suivi des demandes de congés et absences : volumes, types, motifs, statuts et agents absents ce jour.",
    "formation": "Activité de formation : actions, heures, coûts, domaines et agents formés.",
    "departs": "Départs définitifs par motif, taux de départ et projection des départs à la retraite.",
    "prets": "Prêts et avances au personnel : demandes, montants accordés, prêts en cours et mensualités.",
    "mobilite": "Mutations entre directions, demandes de mutation et promotions de grade.",
    "recrutement": "Recrutements par nature, direction et catégorie, et suivi des candidatures et concours.",
    "sante": "Absences pour maladie, maternité et accidents : jours, taux d'absentéisme maladie, indemnités CNAM.",
    "salaires": "Masse salariale, coût employeur, salaires moyens par catégorie, grade et ancienneté.",
    "presence": "Ponctualité et présence : retards, pointages et autorisations de sortie.",
    "budget": "Comparaison du budget de la masse salariale avec le réalisé de la paie.",
    "rendement": "Campagnes d'évaluation : notes de rendement, distribution et lien avec la promotion.",
    "soins": "Couverture médicale : bulletins de soins, frais engagés, remboursements et délais.",
    "social": "Œuvres sociales : aides et dons accordés, primes de scolarité.",
    "endettement": "Cessions sur salaire et taux d'endettement des agents.",
    "services": "Services RH rendus : attestations, demandes de service et intérims.",
}

VIZ_TYPES = {
    "big_number_total": "Indicateur clé",
    "echarts_timeseries_bar": "Barres",
    "echarts_timeseries_line": "Courbe",
    "pie": "Anneau",
    "table": "Tableau",
    "sankey_v2": "Flux (Sankey)",
}

GRAINS = {"P1D": "par jour", "P1M": "par mois", "P3M": "par trimestre", "P1Y": "par année"}

FORMATS = {
    "SMART_NUMBER": "nombre",
    ",.0f": "montant / entier",
    ",.1f": "nombre à 1 décimale",
    ",.2f": "nombre à 2 décimales",
    ".1%": "pourcentage (1 décimale)",
    ".2%": "pourcentage (2 décimales)",
}

SIMPLE_FILTERS = {
    ("DERNIER_12_MOIS", 1): "12 derniers mois",
    ("DOUZE_MOIS", 1): "12 derniers mois",
    ("DERNIER_MOIS", 1): "dernier mois de paie",
    ("EST_ACTIF", 1): "agents actifs",
    ("COD_TYP_BUL", "BN"): "paie mensuelle",
    ("EN_COURS", 1): "en cours",
    ("DERNIERE_SIMULATION", 1): "dernière simulation",
    ("CHANGE_DIRECTION", 1): "changement de direction",
    ("DERNIERE_ANNEE", 1): "dernière année / campagne",
    ("ANNEE_PRECEDENTE", 1): "année / campagne précédente",
    ("ACCORDE", 1): "prêts accordés",
    ("RUBRIQUE_PAIE", 1): "éléments de paie uniquement",
}

SQL_FILTERS = {
    "ANNEE = TO_CHAR(SYSDATE, 'YYYY')": "année en cours",
    "NATURE_ABSENCE <> 'Jours théoriques'": "jours d'absence uniquement",
    "DT_BUL >= ADD_MONTHS(TRUNC(SYSDATE, 'MM'), -36)": "36 derniers mois",
    "DAT_DEM >= ADD_MONTHS(TRUNC(SYSDATE), -12)": "12 derniers mois",
    "DAT_DEBUT >= ADD_MONTHS(TRUNC(SYSDATE), -12) AND DAT_DEBUT <= TRUNC(SYSDATE)": "12 derniers mois",
    "DAT_DEBUT >= ADD_MONTHS(TRUNC(SYSDATE, 'MM'), -24)": "24 derniers mois",
    "MOIS_RESTANTS BETWEEN 0 AND 12": "retraite dans les 12 mois",
    "MOIS_RESTANTS BETWEEN 0 AND 60": "retraite dans les 5 ans",
    "MOIS_RESTANTS BETWEEN 0 AND 120": "retraite dans les 10 ans",
    "MOIS_RESTANTS <= 24": "retraite dans les 24 mois",
    "DAT_DEBUT >= ADD_MONTHS(TRUNC(SYSDATE, 'YYYY'), -48)": "5 dernières années",
    "DAT_DEBUT >= ADD_MONTHS(TRUNC(SYSDATE, 'YYYY'), -108)": "10 dernières années",
    "DAT_EMB >= ADD_MONTHS(TRUNC(SYSDATE, 'YYYY'), -48)": "5 dernières années",
    "DAT_EMB >= ADD_MONTHS(TRUNC(SYSDATE, 'YYYY'), -228)": "20 dernières années",
}

COLUMN_LABELS = {
    "MAT_PERS": "Matricule", "NOM_PRENOM": "Nom et prénom", "DT_BUL": "Date du bulletin", "DT_MOIS": "Mois",
    "NUM_MOIS": "Mois", "DAT_EMB": "Date d'embauche", "DAT_DEM": "Date de demande", "DAT_ACC": "Date de l'accident",
    "DAT_MUT": "Date de mutation", "DAT_GRAD": "Date de promotion", "DAT_DEB_ACTION": "Début de l'action",
    "DAT_FIN_ACTION": "Fin de l'action", "DAT_DEBUT_AUT": "Date de l'autorisation", "DAT_ARRET_TRAV": "Date d'arrêt",
    "DATE_MVT": "Date du mouvement", "NBR_ECH": "Nombre d'échéances", "NBR_JOURS": "Nombre de jours",
    "MNT_ACCORDE": "Montant accordé", "REM_MEN": "Remboursement mensuel", "TAUX_IPP": "Taux d'IPP",
    "FLUX_ORIGINE": "Direction d'origine", "FLUX_DESTINATION": "Direction d'accueil",
    "MOIS_RESTANTS": "Mois restants avant la retraite", "NUM_DECISION": "Numéro de décision",
    "NIVEAU_ESTIME": "Niveau scolaire estimé", "BENEFICIAIRE": "Bénéficiaire",
}

COLUMN_WORDS = {
    "DAT": "date", "DATE": "date", "ANNEE": "année", "AGE": "âge", "ANCIENNETE": "ancienneté",
    "CATEGORIE": "catégorie", "DEBUT": "début", "DEPART": "départ", "ETAT": "état", "DEM": "demande",
    "SOINS": "soins", "ACTE": "acte", "ARRIVEE": "arrivée", "ESTIME": "estimé", "REMUNERATION": "rémunération",
    "PRET": "prêt", "DESTINATION": "destination", "ORIGINE": "origine", "REGION": "région", "RETRAITE": "retraite",
    "RECRUTEMENT": "recrutement", "INTITULE": "intitulé", "PARTICIPANTS": "participants", "COUT": "coût",
    "ACCIDENT": "accident", "CONGE": "congé", "MVT": "mouvement", "RETARD": "retard", "ENDETTEMENT": "endettement",
    "SEMAINE": "semaine", "OBTENU": "obtenu",
}

OPERATORS = {"==": "=", "!=": "≠", ">": ">", "<": "<", ">=": "≥", "<=": "≤", "IN": "dans", "NOT IN": "hors"}

CONVENTIONS = [
    ("Agent actif", "Agent dont l'état de paie est « Actif »."),
    ("12 derniers mois", "Période glissante des 12 derniers mois, par rapport à la date du jour ou au dernier mois "
                         "de paie selon le tableau de bord."),
    ("Dernier mois", "Dernier mois de paie mensuelle disponible."),
    ("Pôle / Direction / Service", "Niveaux de l'organigramme : un pôle regroupe des directions, une direction "
                                   "regroupe des services. Le service est l'affectation réelle de l'agent."),
    ("Effectif payé", "Nombre d'agents distincts ayant reçu un bulletin de paie mensuel."),
    ("Agent-mois", "Un agent payé au cours d'un mois donné (un agent payé 12 mois compte pour 12 agents-mois)."),
    ("Entrée", "Agent payé dans le mois sans l'avoir été le mois précédent."),
    ("Sortie", "Agent payé dans le mois sans l'être le mois suivant."),
    ("Taux de rétention", "Part des agents payés il y a 12 mois qui sont encore payés le dernier mois."),
    ("Turnover", "(Entrées + sorties) / 2, divisé par l'effectif moyen mensuel de la période."),
    ("Responsable", "Agent occupant une fonction de responsabilité : chef de service, chef d'agence, directeur, etc."),
    ("Jours théoriques", "Jours ouvrés attendus : 230 jours par an, soit 230 / 12 par agent payé et par mois."),
    ("Absentéisme toutes causes", "Congés de maladie (y compris maternité, hospitalisation, accident de travail), "
                                  "congés exceptionnels et mises à pied ; hors congé annuel, formations et missions."),
    ("Coût employeur", "Salaire brut + charges patronales (cotisations à la charge de l'employeur)."),
    ("Filtres", "Les filtres d'un dashboard s'appliquent à tous ses graphiques sauf mention contraire. Le filtre "
                "Direction ne propose que les directions du pôle sélectionné."),
]

RULES = {
    ("synthese", "effectif"): "Nombre d'agents distincts ayant reçu une paie mensuelle",
    ("synthese", "entrees"): "Nombre d'agents payés dans le mois qui ne l'étaient pas le mois précédent",
    ("synthese", "sorties"): "Nombre d'agents payés dans le mois qui ne le sont plus le mois suivant",
    ("synthese", "turnover"): "(Entrées + Sorties) / 2, divisé par l'effectif moyen mensuel de la période",
    ("synthese", "retention"): "Agents payés il y a 12 mois et encore payés le dernier mois / agents payés il y a "
                               "12 mois",
    ("synthese", "taux_encadrement"): "Nombre de responsables / effectif payé",
    ("synthese", "responsables"): "Nombre d'agents occupant une fonction de responsabilité (chef de service, chef "
                                  "d'agence, directeur…)",
    ("synthese", "pct_femmes_resp"): "Femmes responsables / total des responsables",
    ("synthese", "masse_brute"): "Somme des salaires bruts mensuels",
    ("synthese", "cout_employeur"): "Salaire brut + charges patronales",
    ("synthese", "cout_moyen"): "Coût employeur total / nombre d'agents-mois payés",
    ("synthese", "jours_absence"): "Somme des jours d'absence : maladie, maternité, hospitalisation, accident de "
                                   "travail, congés exceptionnels et mises à pied",
    ("synthese", "taux_absenteisme"): "Jours d'absence / jours ouvrés théoriques (230 jours par an et par agent payé)",
    ("effectif", "effectif"): "Nombre d'agents distincts",
    ("effectif", "pct_femmes"): "Nombre de femmes / effectif",
    ("effectif", "age_moyen"): "Moyenne de l'âge des agents (en années)",
    ("effectif", "anciennete_moy"): "Moyenne de l'ancienneté des agents depuis leur embauche (en années)",
    ("effectif", "recrut_annee"): "Nombre d'agents embauchés pendant l'année civile en cours",
    ("effectif", "departs_annee"): "Nombre d'agents partis pendant l'année civile en cours",
    ("conges", "jours"): "Somme des jours de congé et d'absence demandés",
    ("conges", "duree_moy"): "Moyenne du nombre de jours par demande",
    ("conges", "solde_moy"): "Moyenne du solde de congé restant par agent (en jours)",
    ("conges", "absents_jour"): "Nombre d'agents dont le congé couvre la date du jour",
    ("formation", "heures"): "Somme des heures des actions de formation",
    ("formation", "cout"): "Somme des coûts TTC des actions de formation (en dinars)",
    ("formation", "cout_moyen"): "Coût total des formations / nombre d'actions",
    ("formation", "participants"): "Somme du nombre de participants de chaque action",
    ("formation", "heures_stagiaire"): "Somme des heures de formation suivies par les participants",
    ("departs", "taux_rotation"): "Départs des 12 derniers mois / effectif actif",
    ("departs", "age_moyen_depart"): "Moyenne de l'âge des agents à leur date de départ",
    ("departs", "retraites"): "Nombre d'agents qui atteindront l'âge de la retraite dans l'horizon indiqué",
    ("prets", "mnt_demande"): "Somme des montants demandés (en dinars)",
    ("prets", "mnt_accorde"): "Somme des montants accordés (en dinars)",
    ("prets", "taux_accord"): "Demandes accordées / total des demandes",
    ("prets", "mnt_moyen"): "Moyenne des montants accordés (en dinars)",
    ("prets", "prets_en_cours"): "Nombre de prêts dont le remboursement n'est pas terminé",
    ("prets", "mensualites"): "Somme des mensualités des prêts en cours de remboursement",
    ("mobilite", "pct_inter_direction"): "Mutations avec changement de direction / total des mutations",
    ("recrutement", "candidats"): "Nombre de candidatures distinctes",
    ("recrutement", "taux_recrutement"): "Candidats recrutés / total des candidats",
    ("recrutement", "part_femmes"): "Nombre de femmes / total",
    ("recrutement", "age_moyen_recrutement"): "Moyenne de l'âge des agents à leur date d'embauche",
    ("sante", "jours_maladie"): "Somme des jours d'arrêt",
    ("sante", "duree_moyenne"): "Moyenne du nombre de jours par arrêt",
    ("sante", "taux_absenteisme"): "Jours d'arrêt maladie / (effectif actif × 230 jours ouvrés par an)",
    ("sante", "demandes_ind"): "Nombre de demandes d'indemnités distinctes",
    ("sante", "jours_ind"): "Somme des jours indemnisés",
    ("salaires", "masse_brute"): "Somme des salaires bruts",
    ("salaires", "cout_employeur"): "Salaire brut + charges patronales",
    ("salaires", "charges_patronales"): "Somme des cotisations à la charge de l'employeur",
    ("salaires", "net_a_payer"): "Somme des nets à payer",
    ("salaires", "salaire_base"): "Somme des salaires de base",
    ("salaires", "indemnites_mensuelles"): "Somme des primes et indemnités mensuelles",
    ("salaires", "retenues_salariales"): "Somme des cotisations sociales retenues sur les salaires",
    ("salaires", "impot"): "Somme de l'impôt sur le revenu et de la contribution sociale de solidarité retenus",
    ("salaires", "salaire_moyen"): "Masse salariale brute / nombre d'agents-mois payés",
    ("salaires", "net_moyen"): "Total des nets à payer / nombre d'agents-mois payés",
    ("salaires", "taux_charges"): "Charges patronales / masse salariale brute",
    ("presence", "minutes_retard"): "Somme des minutes de retard",
    ("presence", "retard_moyen"): "Moyenne des minutes de retard par retard constaté",
    ("presence", "heures_sortie"): "Somme des durées des autorisations de sortie (en heures)",
    ("presence", "agents_pointes"): "Nombre d'agents distincts ayant au moins un pointage",
    ("presence", "presence_moyenne"): "Moyenne des heures de présence par agent et par jour pointé",
    ("presence", "arrivee_moyenne"): "Moyenne de l'heure du premier pointage de la journée",
    ("budget", "budget"): "Somme des montants budgétés",
    ("budget", "realise"): "Somme des montants réalisés en paie",
    ("budget", "ecart"): "Réalisé − budget",
    ("budget", "taux_realisation"): "Réalisé / budget",
    ("rendement", "note_moyenne"): "Moyenne des notes de rendement définitives (sur 20)",
    ("rendement", "part_excellence"): "Nombre de notes supérieures ou égales à 19 / total des notes",
    ("rendement", "taux_promotion"): "Agents promus l'année suivant l'évaluation / agents notés",
    ("rendement", "note_assiduite"): "Moyenne de la note d'assiduité",
    ("rendement", "note_maladie"): "Moyenne de la note liée aux congés de maladie",
    ("soins", "frais_engages"): "Somme des frais de soins engagés par les agents",
    ("soins", "montant_rembourse"): "Somme des montants remboursés",
    ("soins", "taux_remboursement"): "Montant remboursé / frais engagés",
    ("soins", "delai_moyen"): "Moyenne du nombre de jours entre la date des soins et la réception du bulletin",
    ("soins", "rembourse_actes"): "Somme des montants remboursés par acte",
    ("soins", "frais_actes"): "Somme des frais engagés par acte",
    ("social", "montant_accorde"): "Somme des montants accordés",
    ("social", "montant_demande"): "Somme des montants demandés",
    ("social", "taux_acceptation"): "Demandes accordées / total des demandes",
    ("social", "don_moyen"): "Montant total accordé / nombre d'agents bénéficiaires",
    ("social", "primes"): "Nombre de primes de scolarité versées (une par enfant)",
    ("social", "montant_scolarite"): "Somme des primes de scolarité versées",
    ("endettement", "cessions_en_cours"): "Nombre de cessions sur salaire non encore soldées",
    ("endettement", "mensualites_en_cours"): "Somme des retenues mensuelles des cessions en cours",
    ("endettement", "agents_cession"): "Nombre d'agents distincts ayant au moins une cession en cours",
    ("endettement", "taux_moyen"): "Moyenne du taux d'endettement simulé après le nouveau prêt "
                                   "(retenues / salaire)",
    ("endettement", "taux_actuel_moyen"): "Moyenne du taux d'endettement avant le nouveau prêt (retenues / salaire)",
    ("endettement", "agents_seuil"): "Nombre d'agents dont le taux d'endettement atteint ou dépasse 40 %",
    ("services", "delai_moyen"): "Moyenne du nombre de jours entre la demande et la réponse",
    ("services", "en_attente"): "Nombre de demandes à l'état « En attente »",
    ("services", "jours_interim"): "Somme des jours d'intérim",
    ("services", "remplacants"): "Nombre d'agents distincts ayant assuré un intérim",
}


# --------------------------------------------------------------------------- spec reading


def col_label(name: str) -> str:
    if name in COLUMN_LABELS:
        return COLUMN_LABELS[name]
    words = [COLUMN_WORDS.get(w, w.lower()) for w in name.split("_")]
    if words[0] in {"type", "nature", "tranche", "famille", "groupe", "mode", "objet", "statut", "niveau",
                    "modele", "heure", "jour"} and len(words) > 1:
        words.insert(1, "de" if words[1][0] not in "aeiouéâ" else "d'")
    text = " ".join(words).replace("d' ", "d'").replace("modele", "modèle")
    return text[:1].upper() + text[1:]


def fmt_label(fmt: str | None) -> str:
    return FORMATS.get(fmt or "SMART_NUMBER", fmt or "nombre")


def describe_filter(f: dict) -> str:
    if f["expressionType"] == "SQL":
        return SQL_FILTERS.get(f["sqlExpression"], f["sqlExpression"])
    if f["operator"] == "TEMPORAL_RANGE":
        return ""
    known = SIMPLE_FILTERS.get((f["subject"], f["comparator"]))
    if known and f["operator"] == "==":
        return known
    return f"{col_label(f['subject'])} {OPERATORS.get(f['operator'], f['operator'])} {f['comparator']}"


def with_de(text: str) -> str:
    return f"d'{text}" if text[:1] in "aeiouéèêâh" else f"de {text}"


def plain_rule(dashboard: str, metric) -> str | None:
    if (dashboard, metric.name) in RULES:
        return RULES[(dashboard, metric.name)]
    expr = metric.expression.replace(" ", "").upper()
    if expr == "COUNT(*)":
        return f"Nombre {with_de(metric.verbose[:1].lower() + metric.verbose[1:])}"
    if expr == "COUNT(DISTINCTMAT_PERS)":
        return "Nombre d'agents distincts"
    return None


def chart_metrics(params: dict) -> list[str]:
    if "metric" in params:
        return [params["metric"]]
    return list(params.get("metrics", []))


def chart_type(params: dict) -> str:
    viz = params["viz_type"]
    label = VIZ_TYPES.get(viz, viz)
    if viz == "echarts_timeseries_bar":
        if params.get("orientation") == "horizontal":
            label += " horizontales"
        if params.get("stack"):
            label += " empilées"
    if viz == "table":
        label += " agrégé" if params.get("query_mode") == "aggregate" else " détaillé"
    return label


def chart_axes(params: dict) -> str:
    viz = params["viz_type"]
    col = col_label
    if viz == "sankey_v2":
        return f"{col(params['source'])} → {col(params['target'])}"
    if viz == "pie":
        return ", ".join(col(c) for c in params["groupby"])
    if viz == "table":
        cols = params.get("groupby") or params.get("all_columns") or []
        return ", ".join(col(c) for c in cols)
    if viz.startswith("echarts_timeseries"):
        axis = col(params["x_axis"])
        if params.get("x_axis_is_time"):
            axis += f" ({GRAINS.get(params.get('time_grain_sqla'), params.get('time_grain_sqla'))})"
        groups = params.get("groupby") or []
        if groups:
            axis += " ; séries : " + ", ".join(col(g) for g in groups)
        return axis
    return ""


def chart_meaning(params: dict, measures: list[str], filters: list[str]) -> str:
    viz = params["viz_type"]
    col = lambda c: col_label(c).lower()
    what = " et ".join(m.lower() for m in measures) if measures else ""
    if viz == "big_number_total":
        text = f"Valeur unique : {what}"
        if params.get("subheader"):
            text += f" ({params['subheader']})"
    elif viz == "pie":
        text = f"Part de chaque {col(params['groupby'][0])} dans le total : {what}"
    elif viz == "sankey_v2":
        text = f"Flux de {what} de {col(params['source'])} vers {col(params['target'])}"
    elif viz == "table" and params.get("query_mode") != "aggregate":
        text = f"Liste détaillée ({len(params.get('all_columns', []))} colonnes), " \
               f"triée par {col(json_order(params))}"
    elif viz == "table":
        text = f"Tableau de {what} par {', '.join(col(c) for c in params['groupby'])}"
    else:
        grain = GRAINS.get(params.get("time_grain_sqla"), "") if params.get("x_axis_is_time") else ""
        axis = grain if grain else f"par {col(params['x_axis'])}"
        text = f"{what[:1].upper() + what[1:]} {axis}"
        if params.get("groupby"):
            text += ", ventilé par " + ", ".join(col(g) for g in params["groupby"])
        if viz == "echarts_timeseries_line":
            text += " (évolution)"
    if filters:
        text += " — périmètre : " + ", ".join(filters)
    return text[:1].upper() + text[1:] + "."


def json_order(params: dict) -> str:
    order = params.get("order_by_cols") or []
    return json.loads(order[0])[0] if order else ""


class FakeIds(dict):
    def __missing__(self, key):
        return 0


def load_dashboards() -> list[dict]:
    result = []
    for key, label, slug in SUITE:
        spec = importlib.import_module(f"dashboards.{key}")
        metrics = {(ds.key, m.name): m for ds in spec.DATASETS for m in ds.metrics}
        charts = spec.charts(FakeIds())
        order = [k for row in spec.LAYOUT for k, *_ in row]
        order += [k for k in charts if k not in order]
        rows = []
        for ck in order:
            chart = charts[ck]
            p = chart.params
            names = chart_metrics(p)
            measures = [metrics[(chart.dataset, n)].verbose if (chart.dataset, n) in metrics else n for n in names]
            filters = [t for t in (describe_filter(f) for f in p.get("adhoc_filters", [])) if t]
            rows.append({
                "key": ck,
                "name": chart.name,
                "type": chart_type(p),
                "measures": measures,
                "axes": chart_axes(p),
                "filters": filters,
                "meaning": chart_meaning(p, measures, filters),
            })
        filters = []
        for f in spec.FILTERS:
            if isinstance(f, TimeFilter):
                filters.append({"label": f.label, "scope": "Période ; s'applique à : "
                                + ", ".join(charts[c].name for c in f.charts if c in charts)})
            else:
                note = "dépend du filtre Pôle" if f.column == "DIRECTION" and any(
                    getattr(o, "column", None) == "POLE" for o in spec.FILTERS) else ""
                if f.exclude:
                    excl = ", ".join(charts[c].name for c in f.exclude if c in charts)
                    note = (note + " ; " if note else "") + f"sans effet sur : {excl}"
                filters.append({"label": f.label, "scope": note or "tous les graphiques concernés"})
        rules = []
        for m in (m for ds in spec.DATASETS for m in ds.metrics):
            rule = plain_rule(key, m)
            if rule is None:
                print(f"  sans règle de calcul : {key}.{m.name} ({m.expression})")
                rule = m.verbose
            rules.append({"label": m.verbose, "rule": rule, "format": fmt_label(m.d3format)})
        result.append({
            "key": key, "label": label, "title": spec.TITLE, "slug": spec.SLUG,
            "description": DESCRIPTIONS.get(key, ""),
            "rules": rules,
            "charts": rows,
            "filters": filters,
        })
    return result


# --------------------------------------------------------------------------- docx helpers


HEADER_FILL = "1E3A8A"


def set_landscape(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def write_cell(cell, text: str, bold: bool = False, mono: bool = False, color: RGBColor | None = None) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(8 if mono else 9)
    run.bold = bold
    if mono:
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    if color:
        run.font.color.rgb = color


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float],
              mono_cols: set[int] = frozenset()) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    repeat_header(header)
    for i, h in enumerate(headers):
        write_cell(header.cells[i], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(header.cells[i], HEADER_FILL)
    for values in rows:
        cells = table.add_row().cells
        for i, v in enumerate(values):
            write_cell(cells[i], v, mono=i in mono_cols)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph()


# --------------------------------------------------------------------------- document


def public_url(base: str, slug: str, cod_soc: str) -> str:
    url = f"{base.rstrip('/')}/superset/dashboard/{slug}/?standalone=1"
    return f"{url}&cod_soc={quote(cod_soc)}" if cod_soc else url


def build_document(dashboards: list[dict], out: Path, base_url: str, cod_soc: str) -> dict[str, int]:
    doc = Document()
    set_landscape(doc)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    n_charts = sum(len(d["charts"]) for d in dashboards)
    n_rules = sum(len(d["rules"]) for d in dashboards)
    url = {d["key"]: public_url(base_url, d["slug"], cod_soc) for d in dashboards}

    doc.add_heading("Documentation des tableaux de bord GRH", 0)

    doc.add_heading("1. Récapitulatif des dashboards", 1)
    add_table(doc, ["N°", "Dashboard", "URL publique", "Objectif", "Graphiques", "Filtres disponibles"], [
        [str(i), d["title"], url[d["key"]], d["description"], str(len(d["charts"])),
         ", ".join(f["label"] for f in d["filters"])]
        for i, d in enumerate(dashboards, start=1)
    ], [1, 3.5, 6.5, 8, 1.8, 5])

    doc.add_heading("2. Règles de calcul des indicateurs", 1)
    doc.add_paragraph("Chaque indicateur est recalculé à l'affichage, après application des filtres du "
                      "dashboard et du périmètre propre au graphique.")
    add_table(doc, ["Dashboard", "Indicateur", "Règle de calcul", "Format"], [
        [d["title"], r["label"], r["rule"], r["format"]] for d in dashboards for r in d["rules"]
    ], [4, 5.5, 13, 3.3])

    doc.add_heading("3. Détail des graphiques par dashboard", 1)
    for i, d in enumerate(dashboards, start=1):
        doc.add_heading(f"3.{i} {d['title']}", 2)
        p = doc.add_paragraph()
        p.add_run("URL publique : ").bold = True
        p.add_run(url[d["key"]])
        p = doc.add_paragraph()
        p.add_run("Objectif : ").bold = True
        p.add_run(d["description"])
        add_table(doc, ["N°", "Graphique", "Type", "Mesure", "Axe / regroupement", "Périmètre",
                        "Ce que représente le graphique"], [
            [str(n), c["name"], c["type"], ", ".join(c["measures"]) or "Liste détaillée", c["axes"] or "—",
             ", ".join(c["filters"]) or "—", c["meaning"]]
            for n, c in enumerate(d["charts"], start=1)
        ], [0.9, 4.2, 2.6, 3.6, 4, 3.2, 7.6])
        doc.add_paragraph().add_run("Filtres du dashboard").bold = True
        add_table(doc, ["Filtre", "Portée / remarque"], [
            [f["label"], f["scope"]] for f in d["filters"]
        ], [6, 20])

    doc.add_heading("Annexe : définitions", 1)
    add_table(doc, ["Terme", "Définition"], [[t, v] for t, v in CONVENTIONS], [6, 20])

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return {"dashboards": len(dashboards), "charts": n_charts, "rules": n_rules}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--base-url", default="http://localhost:8088", help="adresse publique de Superset")
    parser.add_argument("--cod-soc", default="01", help="société ajoutée aux URL (?cod_soc=..), vide pour aucune")
    args = parser.parse_args()
    dashboards = load_dashboards()
    stats = build_document(dashboards, args.out, args.base_url, args.cod_soc)
    print(", ".join(f"{v} {k}" for k, v in stats.items()))
    print(f"document: {args.out}")


if __name__ == "__main__":
    main()
